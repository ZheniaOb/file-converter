"""Grupa B – tekst i dokumenty: TXT ↔ DOCX (biblioteka python-docx)."""
import re
from pathlib import Path

import docx
from docx.table import Table

from .dane import read_utf8
from .rejestr import ConversionError, converter


@converter(".txt", ".docx")
def txt_to_docx(source: Path, target: Path) -> None:
    """Akapit = blok tekstu oddzielony pustą linią; pojedyncze \\n → złamanie wiersza."""
    document = docx.Document()
    for block in re.split(r"\n\s*\n", read_utf8(source).strip()):
        document.add_paragraph(block)
    document.save(target)


@converter(".docx", ".txt")
def docx_to_txt(source: Path, target: Path) -> None:
    """Akapity i tabele w kolejności z dokumentu; komórki tabeli oddzielone tabulatorem."""
    try:
        document = docx.Document(source)
    except Exception as e:  # python-docx zgłasza różne wyjątki dla uszkodzonego pliku
        raise ConversionError(f"Plik {source.name} nie jest poprawnym dokumentem DOCX ({e}).") from e
    blocks = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            blocks.append("\n".join("\t".join(cell.text for cell in row.cells) for row in item.rows))
        elif item.text.strip():
            blocks.append(item.text)
    target.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
